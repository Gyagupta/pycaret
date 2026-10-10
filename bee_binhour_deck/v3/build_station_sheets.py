"""Per-station Excel tables from the hourly values of the IMD bin-hour analysis.

Input: one row per station per hour of every USED day (the 24 smooth-curve values per day, after gap filling),
as CSV or Parquet, with columns
    station_code, station_name, zone, date, utc_hour, temp_c
  - date      the IMD observation date (YYYY-MM-DD); its 8 readings run 00 UTC (5:30 am IST) to 21 UTC
  - utc_hour  0..23 (the hourly value at that whole UTC hour = hh:30 IST)
  - temp_c    hourly temperature in °C, unrounded
In the analysis notebook this is the table that step 6 ('count hours in 1 °C bins') counts; export it with
    hourly.to_parquet("hourly_values.parquet")   # or .to_csv(...)

Output workbook (one file):
  README_and_Checks     reconciliation against the deck (must all be 0 before the numbers are used)
  <Zone> sheets         each station's hours in bins 24 … 43 and 44+: 10-year count and average year
  All_Degrees_10yr      each station's hours at every 1 °C from 'below 0' up to its highest degree
  All_Degrees_AvgYear   the same, as an average year
  Max_Temperature       highest hourly value per station (and when), and how many stations reach each degree
  Day_Night             per station: day vs night hours and temperatures (average year)
  Hour_Of_Day_by_Zone   average-year hours at 24 °C+ and 35 °C+ for each hour of the day, by zone

Rules (from the deck, slides 14-16): bin = whole degree below the value (24.0-24.9 -> 24), values rounded to
0.000001 °C before binning; 44 °C or more -> '44+'. Average year = for each station and calendar month,
hours ÷ usable days × days in the month (February 28 unless --feb-days says otherwise), 12 months added.
Night = 6:30 pm to 5:30 am IST (UTC hours 13-23 and 0); day = 6:30 am to 5:30 pm IST (UTC hours 1-12).

Usage:
  python -I build_station_sheets.py hourly_values.parquet deck_data.json IMD_BinHours_Stations_v3.xlsx [--feb-days 28]
The script stops (exit 2) if the zone totals do not match Annex 2a/2b exactly, unless --allow-mismatch is given.
"""
import argparse
import json
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ZONES = ["Composite", "Hot-Dry", "Warm-Humid", "Temperate", "Cold", "Undecided"]
BINS = [str(b) for b in range(24, 44)] + ["44+"]
DAYS_IN_MONTH = {1: 31, 2: 28, 3: 31, 4: 30, 5: 31, 6: 30, 7: 31, 8: 31, 9: 30, 10: 31, 11: 30, 12: 31}
NIGHT_UTC = set(range(13, 24)) | {0}

NAVY, BORDER, HIGHLIGHT, LIGHT = "084159", "D6DEE3", "FDEFD9", "F1F4F5"
F_TXT = Font(name="Arial", size=10, color="222222")
F_BOLD = Font(name="Arial", size=10, bold=True, color="222222")
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
F_SUB = Font(name="Arial", size=10, italic=True, color="555555")
F_CALC = Font(name="Arial", size=10, color="000000")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_HI = PatternFill("solid", fgColor=HIGHLIGHT)
FILL_LT = PatternFill("solid", fgColor=LIGHT)
FILL_BAD = PatternFill("solid", fgColor="F8D7DA")
THIN = Side(style="thin", color=BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load(path):
    df = pd.read_parquet(path) if str(path).endswith((".parquet", ".pq")) else pd.read_csv(path, dtype={"station_code": str})
    need = {"station_code", "station_name", "zone", "date", "utc_hour", "temp_c"}
    missing = need - set(df.columns)
    if missing:
        sys.exit(f"input is missing columns: {sorted(missing)}")
    df = df[list(need)].copy()
    df["station_code"] = df["station_code"].astype(str).str.strip()
    df["date"] = pd.to_datetime(df["date"]).dt.normalize()
    df["utc_hour"] = df["utc_hour"].astype(int)
    df["temp_c"] = df["temp_c"].astype(float)
    bad_zone = set(df["zone"]) - set(ZONES)
    if bad_zone:
        sys.exit(f"unknown zone names {sorted(bad_zone)}; expected {ZONES}")
    if df["temp_c"].isna().any():
        sys.exit(f"{int(df['temp_c'].isna().sum())} hourly values are blank; the input must hold used days only")
    dup = df.duplicated(["station_code", "date", "utc_hour"]).sum()
    if dup:
        sys.exit(f"{dup} duplicate station/date/hour rows")
    per_day = df.groupby(["station_code", "date"]).size()
    if (per_day != 24).any():
        sys.exit(f"{int((per_day != 24).sum())} station-days do not have exactly 24 hourly values")
    if df.groupby("station_code")[["station_name", "zone"]].nunique().max().max() > 1:
        sys.exit("a station code maps to more than one name or zone")
    df["degree"] = np.floor(np.round(df["temp_c"].to_numpy(), 6)).astype(int)
    df["month"] = df["date"].dt.month
    df["night"] = df["utc_hour"].isin(NIGHT_UTC)
    return df


def ist_label(utc_hour):
    """Whole UTC hour -> IST clock label, e.g. 0 -> '5:30 am', 13 -> '6:30 pm', 19 -> '12:30 am'."""
    hh = (utc_hour + 5) % 24
    return f"{hh % 12 or 12}:30 {'am' if hh < 12 else 'pm'}"


def bin_label(deg):
    return np.where(deg >= 44, "44+", np.where(deg >= 24, deg.astype(str), "<24"))


def average_year(counts_by_month, days_by_month, feb_days):
    """counts_by_month: DataFrame indexed (station, month) x columns; days_by_month: Series (station, month).
    Returns station x columns, average-year hours."""
    dim = pd.Series({m: (feb_days if m == 2 else v) for m, v in DAYS_IN_MONTH.items()})
    per_day = counts_by_month.div(days_by_month.reindex(counts_by_month.index), axis=0)
    scaled = per_day.mul(dim.reindex(per_day.index.get_level_values("month")).to_numpy(), axis=0)
    return scaled.groupby(level="station_code").sum()


def style_header(ws, row, labels, col=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font, c.fill, c.border = F_HDR, FILL_HDR, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_cell(ws, r, c, v, fmt=None, font=F_TXT, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.border = font, BOX
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("hourly")
    ap.add_argument("deck_data")
    ap.add_argument("out")
    ap.add_argument("--feb-days", type=float, default=28)
    ap.add_argument("--allow-mismatch", action="store_true")
    a = ap.parse_args()

    df = load(a.hourly)
    deck = json.load(open(a.deck_data))
    info = df.groupby("station_code")[["station_name", "zone"]].first()
    days = df.groupby("station_code")["date"].nunique().rename("days_used")
    days_m = df.drop_duplicates(["station_code", "date"]).groupby(["station_code", "month"]).size()
    info = info.join(days)
    info["zone_order"] = info["zone"].map({z: i for i, z in enumerate(ZONES)})
    info = info.sort_values(["zone_order", "station_name"])
    stations = info.index.tolist()

    # ---- counts by bin (10-year) and average year
    df["bin"] = bin_label(df["degree"].to_numpy())
    cnt_bin = df[df["bin"] != "<24"].groupby(["station_code", "bin"]).size().unstack(fill_value=0)
    cnt_bin = cnt_bin.reindex(index=stations, columns=BINS, fill_value=0)
    cnt_bin_m = df[df["bin"] != "<24"].groupby(["station_code", "month", "bin"]).size().unstack(fill_value=0)
    cnt_bin_m = cnt_bin_m.reindex(columns=BINS, fill_value=0)
    avg_bin = average_year(cnt_bin_m, days_m, a.feb_days).reindex(index=stations, columns=BINS, fill_value=0.0)

    # ---- every degree
    lo_all = int(df["degree"].min())
    deg_clip = df["degree"].clip(lower=-1)   # -1 stands for 'below 0'
    deg_cols = list(range(-1, int(df["degree"].max()) + 1))
    cnt_deg = df.assign(d=deg_clip).groupby(["station_code", "d"]).size().unstack(fill_value=0)
    cnt_deg = cnt_deg.reindex(index=stations, columns=deg_cols, fill_value=0)
    cnt_deg_m = df.assign(d=deg_clip).groupby(["station_code", "month", "d"]).size().unstack(fill_value=0)
    cnt_deg_m = cnt_deg_m.reindex(columns=deg_cols, fill_value=0)
    avg_deg = average_year(cnt_deg_m, days_m, a.feb_days).reindex(index=stations, columns=deg_cols, fill_value=0.0)

    # ---- max temperature
    idx = df.groupby("station_code")["temp_c"].idxmax()
    mx = df.loc[idx, ["station_code", "temp_c", "date", "utc_hour"]].set_index("station_code").reindex(stations)
    mn = df.groupby("station_code")["temp_c"].min().reindex(stations)

    # ---- day / night (average year)
    def avg_flag(mask, night):
        sub = df[mask & (df["night"] == night)]
        g = sub.groupby(["station_code", "month"]).size().to_frame("h")
        return average_year(g, days_m, a.feb_days)["h"].reindex(stations, fill_value=0.0)

    dn = pd.DataFrame(index=stations)
    for lab, mask in [("24", df["degree"] >= 24), ("35", df["degree"] >= 35), ("40", df["degree"] >= 40), ("44", df["degree"] >= 44)]:
        dn[f"day_{lab}"] = avg_flag(mask, False)
        dn[f"night_{lab}"] = avg_flag(mask, True)
    tstat = df.groupby(["station_code", "night"])["temp_c"].agg(["mean", "max"]).unstack()
    dn["mean_day"], dn["mean_night"] = tstat[("mean", False)], tstat[("mean", True)]
    dn["max_day"], dn["max_night"] = tstat[("max", False)], tstat[("max", True)]

    # ---- hour of day by zone (average year, simple average of stations)
    hod = {}
    for lab, mask in [("24", df["degree"] >= 24), ("35", df["degree"] >= 35)]:
        g = df[mask].groupby(["station_code", "month", "utc_hour"]).size().unstack(fill_value=0)
        g = g.reindex(columns=range(24), fill_value=0)
        st = average_year(g, days_m, a.feb_days).reindex(stations, fill_value=0.0)
        hod[lab] = st.join(info["zone"]).groupby("zone").mean().reindex(ZONES)

    # ---- reconciliation with the deck (Annex 2a/2b and slide 5)
    ann = {r[0]: r[1:] for r in deck["tables"]["s16_t1"][1:] + deck["tables"]["s17_t1"][1:]}
    hdr = [h.replace("\xa0", " ") for h in deck["tables"]["s16_t1"][0]]
    zone_cnt = cnt_bin.join(info["zone"]).groupby("zone").sum().reindex(ZONES).fillna(0).astype(int)
    checks = []
    for z in ZONES + ["All 129 stations"]:
        j = hdr.index(z) - 1
        ours_st = int((info["zone"] == z).sum()) if z in ZONES else len(stations)
        ours_days = int(info.loc[info["zone"] == z, "days_used"].sum()) if z in ZONES else int(info["days_used"].sum())
        checks.append((z, "Stations", ours_st, int(ann["Stations"][j].replace(",", ""))))
        checks.append((z, "Days used", ours_days, int(ann["Days used"][j].replace(",", ""))))
        for b in BINS:
            ours = int(zone_cnt.loc[z, b]) if z in ZONES else int(zone_cnt[b].sum())
            checks.append((z, b, ours, int(ann[b][j].replace(",", ""))))
    chk = pd.DataFrame(checks, columns=["zone", "item", "this_file", "deck"])
    chk["difference"] = chk["this_file"] - chk["deck"]
    s5 = {r[0]: r[1:] for r in deck["tables"]["s5_t1"][1:]}
    zavg = avg_bin.join(info["zone"]).groupby("zone").mean().reindex(ZONES)
    soft = []
    for z in ZONES:
        ge24 = zavg.loc[z, BINS].sum()
        ge35 = zavg.loc[z, BINS[BINS.index("35"):]].sum()
        ge43 = zavg.loc[z, ["43", "44+"]].sum()
        for lab, v, dk in [("24 °C or more", ge24, s5[z][1]), ("35 °C or more", ge35, s5[z][2]), ("43 °C or more", ge43, s5[z][3]), ("44 °C or more", zavg.loc[z, "44+"], s5[z][4])]:
            dkv = 0.0 if dk.strip().startswith("<") else float(dk.replace(",", ""))
            ok = (v < 1) if dk.strip().startswith("<") else abs(round(v) - dkv) <= 1
            soft.append((z, lab, v, dk, ok))
    b_all = avg_bin.sum(axis=1).mean()
    n44 = int((avg_bin["44+"] >= 1).sum())
    hard_ok = bool((chk["difference"] == 0).all())

    print(f"exact zone checks: {'PASS' if hard_ok else 'FAIL'} ({int((chk['difference'] != 0).sum())} of {len(chk)} differ)")
    print(f"average-year checks vs slide 5: {sum(s[4] for s in soft)} of {len(soft)} within rounding")
    print(f"B (all stations, average year, 24 °C+): {b_all:,.1f} (deck 6,372); stations with ≥1 h a year at 44 °C+: {n44} (deck 50)")
    if not hard_ok and not a.allow_mismatch:
        print(chk[chk["difference"] != 0].to_string(index=False))
        sys.exit(2)

    # ---------------------------------------------------------------- workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "README_and_Checks"
    ws["A1"] = "IMD bin-hours by station (v3): checks against the deck"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Built from the hourly values of the analysis. Bin rule, average year and day/night as in the deck. "
                "Exact checks must all be 0 before these tables are used.")
    ws["A2"].font = F_SUB
    ws["A4"] = f"Exact zone checks vs Annex 2a/2b: {'PASS' if hard_ok else 'FAIL'}"
    ws["A4"].font = F_BOLD
    ws["A4"].fill = FILL_HI if hard_ok else FILL_BAD
    ws["A5"] = (f"Stations {len(stations)}; days used {int(info['days_used'].sum()):,}; hourly values {len(df):,}; "
                f"lowest {df['temp_c'].min():.1f} °C; highest {df['temp_c'].max():.1f} °C; February days in the average year: {a.feb_days:g}")
    ws["A5"].font = F_TXT
    style_header(ws, 7, ["Zone", "Item", "This file", "Deck (Annex 2a/2b)", "Difference"])
    for i, row in enumerate(chk.itertuples(index=False)):
        r = 8 + i
        for j, v in enumerate(row):
            write_cell(ws, r, 1 + j, v, "#,##0" if j >= 2 else None, fill=FILL_BAD if (j == 4 and v != 0) else None)
    r = 8 + len(chk) + 2
    style_header(ws, r, ["Zone", "Average year", "This file", "Deck (slide 5, rounded)", "Within rounding"])
    for i, (z, lab, v, dk, ok) in enumerate(soft):
        for j, val in enumerate([z, lab, v, dk, "yes" if ok else "NO"]):
            write_cell(ws, r + 1 + i, 1 + j, val, "#,##0.0" if j == 2 else None, fill=None if ok or j != 4 else FILL_BAD)
    r = r + len(soft) + 2
    write_cell(ws, r, 1, "B: average of all stations, hours at 24 °C+ a year", font=F_BOLD)
    write_cell(ws, r, 3, round(b_all, 1), "#,##0.0")
    write_cell(ws, r, 4, "6,372")
    write_cell(ws, r + 1, 1, "Stations with at least 1 hour a year at 44 °C+", font=F_BOLD)
    write_cell(ws, r + 1, 3, n44, "0")
    write_cell(ws, r + 1, 4, "50")
    for col, w in zip("ABCDE", (34, 22, 16, 22, 14)):
        ws.column_dimensions[col].width = w

    # zone sheets
    for z in ZONES:
        sts = [s for s in stations if info.loc[s, "zone"] == z]
        if not sts:
            continue
        ws = wb.create_sheet(z[:31])
        ws["A1"] = f"{z}: hours in each 1 °C bin, by station"
        ws["A1"].font = F_TITLE
        ws["A2"] = "Block 1: hours counted 2016–2025 (all usable days). Block 2: average year (Jan–Dec). Bin 24 = 24.0–24.9 °C."
        ws["A2"].font = F_SUB
        cols = ["Station code", "Station", "Days used"] + BINS + ["24 °C and above", "35 °C and above"]
        for block, (lab, data, fmt) in enumerate([("Hours counted 2016–2025", cnt_bin, "#,##0"), ("Average year", avg_bin, "#,##0.0")]):
            top = 4 if block == 0 else 4 + len(sts) + 6
            ws.cell(row=top, column=1, value=lab).font = F_BOLD
            style_header(ws, top + 1, cols)
            for i, s in enumerate(sts):
                r = top + 2 + i
                write_cell(ws, r, 1, s)
                write_cell(ws, r, 2, info.loc[s, "station_name"])
                write_cell(ws, r, 3, int(info.loc[s, "days_used"]), "#,##0")
                for j, b in enumerate(BINS):
                    v = data.loc[s, b]
                    write_cell(ws, r, 4 + j, int(v) if block == 0 else float(v), fmt)
                c0, c44, c35 = get_column_letter(4), get_column_letter(3 + len(BINS)), get_column_letter(4 + BINS.index("35"))
                write_cell(ws, r, 4 + len(BINS), f"=SUM({c0}{r}:{c44}{r})", fmt, F_CALC, FILL_LT)
                write_cell(ws, r, 5 + len(BINS), f"=SUM({c35}{r}:{c44}{r})", fmt, F_CALC, FILL_LT)
            r = top + 2 + len(sts)
            write_cell(ws, r, 2, "Zone total" if block == 0 else "Zone average (each station equal)", font=F_BOLD, fill=FILL_HI)
            fn = "SUM" if block == 0 else "AVERAGE"
            for j in range(3 if block == 0 else 4, 6 + len(BINS)):
                L = get_column_letter(j)
                write_cell(ws, r, j, f"={fn}({L}{top + 2}:{L}{r - 1})", fmt, F_CALC, FILL_HI)
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 28
        for j in range(3, 6 + len(BINS)):
            ws.column_dimensions[get_column_letter(j)].width = 9 if j < 4 + len(BINS) else 12
        ws.freeze_panes = "D6"

    # every degree
    for name, data, fmt, sub in [("All_Degrees_10yr", cnt_deg, "#,##0", "Hours counted 2016–2025 over all usable days."),
                                 ("All_Degrees_AvgYear", avg_deg, "#,##0.0", "Average year (January–December).")]:
        ws = wb.create_sheet(name)
        ws["A1"] = "Hours at every degree, by station"
        ws["A1"].font = F_TITLE
        ws["A2"] = (f"{sub} Column 25 = 25.0–25.9 °C. 'Below 0' holds everything under 0 °C (lowest value {lo_all} °C bin). "
                    "Columns run to the highest degree reached at any station.")
        ws["A2"].font = F_SUB
        labels = ["Station code", "Station", "Zone", "Highest °C"] + ["Below 0" if dgr == -1 else dgr for dgr in deg_cols] + ["All hours"]
        style_header(ws, 4, labels)
        for i, s in enumerate(stations):
            r = 5 + i
            write_cell(ws, r, 1, s)
            write_cell(ws, r, 2, info.loc[s, "station_name"])
            write_cell(ws, r, 3, info.loc[s, "zone"])
            write_cell(ws, r, 4, round(float(mx.loc[s, "temp_c"]), 1), "0.0")
            for j, dgr in enumerate(deg_cols):
                v = data.loc[s, dgr]
                c = write_cell(ws, r, 5 + j, int(v) if fmt == "#,##0" else float(v), fmt)
                if dgr >= 44 and v > 0:
                    c.fill = FILL_HI
            L0, L1 = get_column_letter(5), get_column_letter(4 + len(deg_cols))
            write_cell(ws, r, 5 + len(deg_cols), f"=SUM({L0}{r}:{L1}{r})", fmt, F_CALC, FILL_LT)
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 12
        ws.freeze_panes = "E5"

    # max temperature
    ws = wb.create_sheet("Max_Temperature")
    ws["A1"] = "Highest temperature reached at each station, and how many stations reach each degree"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Highest hourly value on the smooth curve, 2016–2025 (usable days only). IST time = UTC + 5:30."
    ws["A2"].font = F_SUB
    style_header(ws, 4, ["Station code", "Station", "Zone", "Highest °C", "Date", "Time (IST)", "Lowest °C",
                         "Hours a year at 44 °C+", "Hours a year at 40 °C+"])
    order = mx.sort_values("temp_c", ascending=False).index
    for i, s in enumerate(order):
        r = 5 + i
        ist = ist_label(int(mx.loc[s, "utc_hour"]))
        vals = [s, info.loc[s, "station_name"], info.loc[s, "zone"], round(float(mx.loc[s, "temp_c"]), 2),
                mx.loc[s, "date"].strftime("%Y-%m-%d"), ist, round(float(mn.loc[s]), 2),
                float(avg_bin.loc[s, "44+"]), float(avg_deg.loc[s, [c for c in deg_cols if c >= 40]].sum())]
        for j, v in enumerate(vals):
            write_cell(ws, r, 1 + j, v, {3: "0.00", 6: "0.00", 7: "#,##0.0", 8: "#,##0.0"}.get(j))
    r0 = 5 + len(order) + 2
    style_header(ws, r0, ["Degree reached", "Stations reaching it at least once in 10 years",
                          "Stations with at least 1 hour a year (average year)", "Of all stations (at least once)"])
    top_deg = int(df["degree"].max())
    for k, dgr in enumerate(range(35, top_deg + 1)):
        r = r0 + 1 + k
        once = int((mx["temp_c"].apply(lambda t: np.floor(np.round(t, 6))) >= dgr).sum())
        yearly = int((avg_deg[[c for c in deg_cols if c >= dgr]].sum(axis=1) >= 1).sum())
        write_cell(ws, r, 1, f"{dgr} °C or more", font=F_BOLD, fill=FILL_HI if dgr in (40, 44, 45) else None)
        write_cell(ws, r, 2, once, "0")
        write_cell(ws, r, 3, yearly, "0")
        write_cell(ws, r, 4, f"=B{r}/{len(stations)}", "0.0%", F_CALC)
    for col, w in zip("ABCDEFGHI", (12, 28, 12, 11, 12, 11, 11, 14, 14)):
        ws.column_dimensions[col].width = w

    # day / night
    ws = wb.create_sheet("Day_Night")
    ws["A1"] = "Day and night, by station (average year)"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Day = 6:30 am to 5:30 pm IST; night = 6:30 pm to 5:30 am IST (12 hourly values each)."
    ws["A2"].font = F_SUB
    labels = ["Station code", "Station", "Zone",
              "24 °C+ day", "24 °C+ night", "Night share of 24 °C+",
              "35 °C+ day", "35 °C+ night", "Night share of 35 °C+",
              "40 °C+ day", "40 °C+ night", "44 °C+ day", "44 °C+ night",
              "Mean day °C", "Mean night °C", "Highest day °C", "Highest night °C"]
    style_header(ws, 4, labels)
    for i, s in enumerate(stations):
        r = 5 + i
        row = [s, info.loc[s, "station_name"], info.loc[s, "zone"],
               dn.loc[s, "day_24"], dn.loc[s, "night_24"], f"=IF(D{r}+E{r}=0,0,E{r}/(D{r}+E{r}))",
               dn.loc[s, "day_35"], dn.loc[s, "night_35"], f"=IF(G{r}+H{r}=0,0,H{r}/(G{r}+H{r}))",
               dn.loc[s, "day_40"], dn.loc[s, "night_40"], dn.loc[s, "day_44"], dn.loc[s, "night_44"],
               dn.loc[s, "mean_day"], dn.loc[s, "mean_night"], dn.loc[s, "max_day"], dn.loc[s, "max_night"]]
        for j, v in enumerate(row):
            fmt = "0.0%" if j in (5, 8) else ("0.0" if j >= 13 else ("#,##0.0" if j >= 3 else None))
            write_cell(ws, r, 1 + j, float(v) if isinstance(v, (np.floating, float)) else v, fmt, F_CALC if j in (5, 8) else F_TXT)
    ws.column_dimensions["B"].width = 28
    ws.freeze_panes = "D5"

    # hour of day by zone
    ws = wb.create_sheet("Hour_Of_Day_by_Zone")
    ws["A1"] = "Average-year hours at each hour of the day, by zone (each station equal)"
    ws["A1"].font = F_TITLE
    for k, lab in enumerate(["24", "35"]):
        top = 3 + k * 30
        ws.cell(row=top, column=1, value=f"Hours a year at {lab} °C or more").font = F_BOLD
        style_header(ws, top + 1, ["Hour (IST)", "Day or night"] + ZONES)
        for i, h in enumerate(list(range(0, 24))):
            r = top + 2 + i
            write_cell(ws, r, 1, ist_label(h))
            write_cell(ws, r, 2, "Night" if h in NIGHT_UTC else "Day")
            for j, z in enumerate(ZONES):
                write_cell(ws, r, 3 + j, float(hod[lab].loc[z, h]) if z in hod[lab].index and not np.isnan(hod[lab].loc[z, h]) else 0.0, "#,##0.0")

    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = False
    wb.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
