"""Build IMD_BinHours_Zone_National_v3.xlsx from the numbers in IMD_BinHours_CLASP_FINAL_v2.pptx.

Every hard-coded number is copied from the deck (deck_data.json, extracted by extract_deck_data.py)
and shown in blue; every total, band, share and check is an Excel formula (black, or green when it
links to another sheet). Usage: python -I build_excel.py deck_data.json out.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DATA, OUT = sys.argv[1], sys.argv[2]
d = json.load(open(DATA))

NAVY, BORDER, LIGHT, HIGHLIGHT = "084159", "D6DEE3", "F1F4F5", "FDEFD9"
F_IN = Font(name="Arial", size=10, color="0000FF")          # copied from the deck
F_CALC = Font(name="Arial", size=10, color="000000")        # formula on this sheet
F_LINK = Font(name="Arial", size=10, color="008000")        # formula linking another sheet
F_TXT = Font(name="Arial", size=10, color="222222")
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
F_SUB = Font(name="Arial", size=10, italic=True, color="555555")
F_BOLD = Font(name="Arial", size=10, bold=True, color="222222")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_HI = PatternFill("solid", fgColor=HIGHLIGHT)
FILL_LT = PatternFill("solid", fgColor=LIGHT)
THIN = Side(style="thin", color=BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

BINS = [str(b) for b in range(24, 44)] + ["44+"]
ZONES = ["Composite", "Hot-Dry", "Warm-Humid", "Temperate", "Cold", "Undecided"]


def num(v):
    """Deck cell text -> number. '–', '-', '' and 'not counted' -> None. '<1' and words kept as text."""
    s = str(v).strip().replace(",", "").replace("\xa0", "")
    if s in ("–", "-", "", "not counted", "None"):
        return None
    if s.startswith("<"):
        return s
    try:
        return float(s) if "." in s else int(s)
    except ValueError:
        return str(v).strip()


def table(key):
    return d["tables"][key]


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB


def header(ws, row, labels, col=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font, c.fill, c.alignment, c.border = F_HDR, FILL_HDR, CENTER, BOX


def put(ws, row, col, value, font=F_IN, fmt=None, fill=None, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Arial", size=10, color=font.color, bold=bold)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if isinstance(value, str) and not value.startswith("="):
        c.alignment = Alignment(wrap_text=True, vertical="center")
    return c


def note(ws, row, text, col=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_SUB
    return c


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
title(ws, "IMD bin-hours: zone and national tables (v3)",
      "All numbers are taken from IMD_BinHours_CLASP_FINAL_v2.pptx (CLASP analysis of IMD 2016–2025 data, 129 stations). Slide numbers refer to the v3 deck.")
rows = [
    ("Sheet", "What it shows", "Source in the deck"),
    ("Proposed_vs_Current", "Bin-by-bin national table: today's IS 18154 table vs proposed A (by population) and B (each station equal), "
     "hours a year and of 1,600; bands; side-by-side summary of what changes.", "Slides 4, 6, 7 and 17 (Annex 1)"),
    ("Population_Weighting", "Share of stations vs share of people by zone; why A differs from B; check that B = station-weighted "
     "average of zone averages.", "Slide 6 table and speaker notes; slide 5"),
    ("Zone_Bins_10yr", "Hours counted in each 1 °C bin (24…43, 44+) by zone over 2016–2025, all usable days (10-year totals, "
     "not scaled to a year), with sub-totals and each bin's share.", "Slides 18–19 (Annex 2a, 2b)"),
    ("Zone_Average_Year", "Average hours in ONE year at or above 24, 35, 43 and 44 °C by zone and for India.", "Slide 5"),
    ("Hours_40C_and_Above", "Hours at 40, 41, 42, 43 and 44+ °C: national average year and 10-year zone counts; stations "
     "reaching 44 °C; the 8 stations with most hours at 44 °C+.", "Slides 9, 10, 17, 19"),
    ("Day_Night", "Hours at 24–34 °C and 35 °C+ by hour of the day and by month (India, by population, average year); "
     "day and night totals and shares.", "Slide 8 charts"),
]
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = ws.cell(row=4 + i, column=1 + j, value=v)
        c.border, c.alignment = BOX, WRAP
        c.font = F_HDR if i == 0 else F_TXT
        if i == 0:
            c.fill = FILL_HDR
r = 4 + len(rows) + 1
for line in [
    "Colour code: blue = number copied from the deck; black = formula on the same sheet; green = formula linking another sheet.",
    "Rounding: the deck shows rounded values. Totals stated in the deck come from unrounded numbers, so a SUM of the rounded "
    "bins can differ by a few hours; both are shown and labelled.",
    "Bin rule (deck slide 15): bin 24 = 24.0–24.9 °C; 35.4 °C → bin 35; 44 °C or more → bin 44+.",
    "Average year = for each station, month by month, hours per usable day × days in the month, 12 months added (not a 10-year total).",
    "NOT IN THIS FILE (needs station-level data, which the deck does not carry): per-station bin tables by zone, per-station "
    "hours at every degree from the lowest to the highest temperature, each station's maximum temperature, and per-station "
    "day/night hours. These are produced by build_station_sheets.py from the analysis output.",
]:
    c = ws.cell(row=r, column=1, value=line)
    c.font, c.alignment = F_TXT, WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 30
    r += 1
widths(ws, [24, 90, 34])

# ---------------------------------------------------------------- Proposed_vs_Current
ws = wb.create_sheet("Proposed_vs_Current")
title(ws, "Proposed bin-hour table vs today's table (IS 18154)",
      "Hours a year = average year, January to December. Of 1,600 = the same hours scaled so the bins add to 1,600, as in BEE's table.")
ann = {r[0]: [num(x) for x in r[1:]] for r in table("s15_t1")[1:] + table("s15_t2")[1:] if r[0] != "Bin °C"}
# deck column order: A hrs, A of1600, B hrs, B of1600, Today hrs, Today of1600 -> sheet order Today, A, B
header(ws, 4, ["Bin °C", "Today: hours a year", "Today: of 1,600", "A: hours a year", "A: of 1,600",
               "B: hours a year", "B: of 1,600", "A − Today (of 1,600)", "B − Today (of 1,600)"])
FIRST = 5
for i, b in enumerate(BINS):
    row = FIRST + i
    a_h, a_s, b_h, b_s, t_h, t_s = ann[b]
    put(ws, row, 1, b if b == "44+" else int(b), F_IN)
    for col, v, fmt in [(2, t_h, "#,##0"), (3, t_s, "0.0"), (4, a_h, "#,##0"), (5, a_s, "0.0"), (6, b_h, "#,##0"), (7, b_s, "0.0")]:
        put(ws, row, col, v if v is not None else "not counted", F_IN, fmt)
    put(ws, row, 8, f"=E{row}-N(C{row})", F_CALC, "+0.0;-0.0;0.0")
    put(ws, row, 9, f"=G{row}-N(C{row})", F_CALC, "+0.0;-0.0;0.0")
LAST = FIRST + len(BINS) - 1                       # row of 44+
R35 = FIRST + BINS.index("35")
r = LAST + 1
put(ws, r, 1, "Sum of the rounded bins", F_TXT, bold=True)
for col in range(2, 8):
    L = get_column_letter(col)
    put(ws, r, col, f"=SUM({L}{FIRST}:{L}{LAST})", F_CALC, "#,##0.0" if col % 2 else "#,##0")
SUMROW = r
r += 1
put(ws, r, 1, "Total stated in the deck (from unrounded numbers)", F_TXT, bold=True)
tot = ann["Total"]
for col, v, fmt in [(2, tot[4], "#,##0"), (3, tot[5], "#,##0"), (4, tot[0], "#,##0"), (5, tot[1], "#,##0"), (6, tot[2], "#,##0"), (7, tot[3], "#,##0")]:
    put(ws, r, col, v, F_IN, fmt)
r += 1
put(ws, r, 1, "35 °C and above: sum of the rounded bins", F_TXT, bold=True)
for col in range(2, 8):
    L = get_column_letter(col)
    put(ws, r, col, f"=SUM({L}{R35}:{L}{LAST})", F_CALC, "#,##0.0" if col % 2 else "#,##0", FILL_HI)
put(ws, r, 8, f"=E{r}-C{r}", F_CALC, "+0.0;-0.0;0.0", FILL_HI)
put(ws, r, 9, f"=G{r}-C{r}", F_CALC, "+0.0;-0.0;0.0", FILL_HI)
r += 1
put(ws, r, 1, "35 °C and above: stated in the deck", F_TXT, bold=True)
a35 = ann["35 and above"]
for col, v, fmt in [(2, a35[4], "#,##0"), (3, a35[5], "0.0"), (4, a35[0], "#,##0"), (5, a35[1], "0.0"), (6, a35[2], "#,##0"), (7, a35[3], "0.0")]:
    put(ws, r, col, v, F_IN, fmt)
note(ws, r + 1, "Today has no 44 °C+ bin ('not counted'); the difference columns treat it as 0. Deck: Annex 1 (slide 17).")

# bands
r += 3
BAND_ROW0 = r
header(ws, r, ["Band (of 1,600)", "Today", "A: by population", "B: each station equal", "A − Today", "B − Today"])
bands = [("24–29 °C", "24", "29"), ("30–34 °C", "30", "34"), ("35–39 °C", "35", "39"), ("40–43 °C", "40", "43"), ("44 °C and above", "44+", "44+")]
for k, (lab, lo, hi) in enumerate(bands):
    rr = r + 1 + k
    r1, r2 = FIRST + BINS.index(lo), FIRST + BINS.index(hi)
    put(ws, rr, 1, lab, F_TXT, bold=True)
    put(ws, rr, 2, f"=SUM(C{r1}:C{r2})", F_CALC, "0.0")
    put(ws, rr, 3, f"=SUM(E{r1}:E{r2})", F_CALC, "0.0")
    put(ws, rr, 4, f"=SUM(G{r1}:G{r2})", F_CALC, "0.0")
    put(ws, rr, 5, f"=C{rr}-B{rr}", F_CALC, "+0.0;-0.0;0.0")
    put(ws, rr, 6, f"=D{rr}-B{rr}", F_CALC, "+0.0;-0.0;0.0")
rr = r + 1 + len(bands)
put(ws, rr, 1, "All bins", F_TXT, bold=True)
for col in range(2, 5):
    L = get_column_letter(col)
    put(ws, rr, col, f"=SUM({L}{r + 1}:{L}{rr - 1})", F_CALC, "0.0")
note(ws, rr + 1, "Bands are sums of the rounded bins above, so they may differ by ±0.2 from the deck's unrounded totals.")

# side-by-side summary
r = rr + 3
SUMMARY_ROW0 = r
header(ws, r, ["What", "Today (IS 18154)", "Proposed (IMD 2016–2025)"])
summary = [
    ("Data behind the table", "54 cities; years and method not published", "129 IMD stations with full-day readings, 2016–2025; method documented"),
    ("How hours are counted", "Not published", "8 readings a day → smooth curve → 24 hourly values"),
    ("Bins", "1 °C bins, 24 to 43 °C", "1 °C bins, 24 to 43 °C, plus '44 °C and above'"),
    ("Day and night", "Counted (5,774 hours a year is more than a 12-hour day can hold)", "Counted, all 24 hours"),
    ("How places are combined", "Not published", "A: by population (state 2026 projection, shared by district); B: each station equal"),
    ("Zone tables", "None", "6 zones plus India"),
    ("Hours at 24 °C or more, a year", "=TEXT(B{0},\"#,##0\")".format(SUMROW + 1), '="A "&TEXT(D{0},"#,##0")&"  ·  B "&TEXT(F{0},"#,##0")'.format(SUMROW + 1)),
    ("35 °C or more, of 1,600", "=TEXT(C{0},\"0\")".format(SUMROW + 3), '="A "&TEXT(E{0},"0")&"  ·  B "&TEXT(G{0},"0")'.format(SUMROW + 3)),
    ("Top bin, of 1,600", '="43 °C: "&TEXT(C{0},"0")'.format(FIRST + BINS.index("43")),
     '="43 °C: A "&TEXT(E{0},"0.0")&" · B "&TEXT(G{0},"0.0")&";  44 °C+: A "&TEXT(E{1},"0.0")&" · B "&TEXT(G{1},"0.0")'.format(FIRST + BINS.index("43"), LAST)),
]
for k, (a, b, c) in enumerate(summary):
    rr = r + 1 + k
    put(ws, rr, 1, a, F_TXT, bold=True)
    put(ws, rr, 2, b, F_CALC if str(b).startswith("=") else F_IN)
    put(ws, rr, 3, c, F_CALC if str(c).startswith("=") else F_IN)
    ws.row_dimensions[rr].height = 30
note(ws, rr + 1, "Sources: deck slides 4, 6, 7, 8, 14 and 17; IS 18154 (Part 1):2023 Table 4A for today's table.")
widths(ws, [44, 30, 46, 16, 14, 16, 14, 18, 18])
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Population_Weighting
ws = wb.create_sheet("Population_Weighting")
title(ws, "Where people live vs where the stations are",
      "A weights each station by the people it stands for; B counts every station once. Shares from the deck (slide 6 notes).")
header(ws, 4, ["Zone", "Stations used", "Share of stations (calculated)", "Share of stations (deck)", "Share of people (deck)",
               "People share ÷ station share", "Average hours a year at 24 °C+", "Average hours a year at 35 °C+"])
notes = d["notes"]["s6"]
shares = {}
for line in notes.splitlines():
    parts = line.split()
    if len(parts) >= 3 and parts[-1].endswith("%") and parts[-2].endswith("%"):
        shares[" ".join(parts[:-2])] = (float(parts[-2][:-1]) / 100, float(parts[-1][:-1]) / 100)
z5 = {r[0]: r for r in table("s5_t1")[1:]}
order = ["Composite", "Warm-Humid", "Hot-Dry", "Undecided", "Temperate", "Cold"]
for k, z in enumerate(order):
    rr = 5 + k
    put(ws, rr, 1, z, F_TXT, bold=True)
    put(ws, rr, 2, num(z5[z][1]), F_IN, "0")
    put(ws, rr, 3, f"=B{rr}/$B$11", F_CALC, "0.0%")
    put(ws, rr, 4, shares[z][0], F_IN, "0.0%")
    put(ws, rr, 5, shares[z][1], F_IN, "0.0%")
    put(ws, rr, 6, f"=E{rr}/D{rr}", F_CALC, "0.00")
    put(ws, rr, 7, num(z5[z][2]), F_IN, "#,##0")
    put(ws, rr, 8, num(z5[z][3]), F_IN, "#,##0")
put(ws, 11, 1, "Total", F_TXT, bold=True)
for col, fmt in [(2, "0"), (3, "0.0%"), (4, "0.0%"), (5, "0.0%")]:
    L = get_column_letter(col)
    put(ws, 11, col, f"=SUM({L}5:{L}10)", F_CALC, fmt)
header(ws, 13, ["Check", "Calculated", "Deck"])
put(ws, 14, 1, "B, hours at 24 °C+ = station-weighted average of zone averages", F_TXT, bold=True)
put(ws, 14, 2, "=SUMPRODUCT(B5:B10,G5:G10)/B11", F_CALC, "#,##0.0")
put(ws, 14, 3, num(table("s6_t1")[1][2]), F_IN, "#,##0")
put(ws, 15, 1, "B, hours at 35 °C+ = station-weighted average of zone averages", F_TXT, bold=True)
put(ws, 15, 2, "=SUMPRODUCT(B5:B10,H5:H10)/B11", F_CALC, "#,##0.0")
put(ws, 15, 3, num(table("s6_t1")[2][2]), F_IN, "#,##0")
note(ws, 16, "A cannot be rebuilt from zone shares: it weights each station by its own district population, not by zone.")
header(ws, 18, ["Result (slide 6)", "A: by population", "B: each station equal", "Today (IS 18154)"])
for k, row in enumerate(table("s6_t1")[1:]):
    rr = 19 + k
    put(ws, rr, 1, row[0], F_TXT, bold=True)
    for j in range(3):
        v = num(row[1 + j])
        put(ws, rr, 2 + j, v, F_IN, "#,##0.0" if isinstance(v, float) else "#,##0")
note(ws, 19 + len(table("s6_t1")) - 1,
     "A: each state's projected 2026 population (National Commission on Population, MoHFW, 2019) shared among its stations "
     "by Census 2011 district population. B: simple average of the 129 stations.")
widths(ws, [58, 16, 18, 18, 18, 18, 20, 20])

# ---------------------------------------------------------------- Zone_Bins_10yr
ws = wb.create_sheet("Zone_Bins_10yr")
title(ws, "Hours counted in each 1 °C bin, by zone, 2016–2025 (10-year totals)",
      "All usable days (all 8 readings, or 1–2 filled). Not scaled to a year: stations have different numbers of usable days.")
z = {r[0]: [num(x) for x in r[1:]] for r in table("s16_t1")[1:] + table("s17_t1")[1:]}
header(ws, 4, ["Bin °C"] + ZONES + ["All 129 stations (sum of zones)", "All 129 stations (deck)", "Check (should be 0)"])
labels = ["Stations", "Days used"] + BINS
for k, lab in enumerate(labels):
    rr = 5 + k
    put(ws, rr, 1, lab if lab in ("Stations", "Days used", "44+") else int(lab), F_TXT if k < 2 else F_IN, bold=k < 2)
    for j in range(6):
        put(ws, rr, 2 + j, z[lab][j], F_IN, "#,##0")
    put(ws, rr, 8, f"=SUM(B{rr}:G{rr})", F_CALC, "#,##0")
    put(ws, rr, 9, z[lab][6], F_IN, "#,##0")
    put(ws, rr, 10, f"=H{rr}-I{rr}", F_CALC, "#,##0")
B0, B1 = 7, 7 + len(BINS) - 1          # bin rows 24..44+
rr = B1 + 1
put(ws, rr, 1, "Total (sum of bins)", F_TXT, bold=True)
for col in range(2, 9):
    L = get_column_letter(col)
    put(ws, rr, col, f"=SUM({L}{B0}:{L}{B1})", F_CALC, "#,##0", FILL_LT)
TOT = rr
rr += 1
put(ws, rr, 1, "Total stated in the deck", F_TXT, bold=True)
for j in range(6):
    put(ws, rr, 2 + j, z["Total"][j], F_IN, "#,##0")
put(ws, rr, 8, z["Total"][6], F_IN, "#,##0")
rr += 1
put(ws, rr, 1, "Check (should be 0)", F_TXT, bold=True)
for col in range(2, 9):
    L = get_column_letter(col)
    put(ws, rr, col, f"={L}{TOT}-{L}{TOT + 1}", F_CALC, "#,##0")
rr += 2
header(ws, rr, ["Sub-total"] + ZONES + ["All 129 stations"])
for k, (lab, lo) in enumerate([("35 °C and above", "35"), ("40 °C and above", "40"), ("43 °C and above", "43"), ("44 °C and above", "44+")]):
    r2 = rr + 1 + k
    r1 = B0 + BINS.index(lo)
    put(ws, r2, 1, lab, F_TXT, bold=True)
    for col in range(2, 9):
        L = get_column_letter(col)
        put(ws, r2, col, f"=SUM({L}{r1}:{L}{B1})", F_CALC, "#,##0")
rr = r2 + 2
header(ws, rr, ["Share of the zone's counted hours"] + ZONES + ["All 129 stations"])
for k, b in enumerate(BINS):
    r2 = rr + 1 + k
    put(ws, r2, 1, b if b == "44+" else int(b), F_TXT)
    for col in range(2, 9):
        L = get_column_letter(col)
        put(ws, r2, col, f"={L}{B0 + k}/{L}${TOT}", F_CALC, "0.00%")
note(ws, r2 + 1, "Source: deck slides 18–19 (Annex 2a, 2b). Compare zones with the average-year sheet, not with these totals.")
widths(ws, [32, 13, 13, 13, 13, 13, 13, 18, 16, 14])
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Zone_Average_Year
ws = wb.create_sheet("Zone_Average_Year")
title(ws, "Average hours in ONE year at or above each temperature, by zone",
      "Average year built from 10 years of IMD data (2016–2025), January to December. Each zone = simple average of its stations.")
t5 = table("s5_t1")
header(ws, 4, t5[0])
for k, row in enumerate(t5[1:]):
    rr = 5 + k
    hi = FILL_HI if row[0].startswith("All India") or row[0].startswith("Today") else None
    put(ws, rr, 1, row[0], F_TXT, bold=True, fill=hi)
    for j, v in enumerate(row[1:]):
        val = num(v)
        if val is None:
            val = v
        c = put(ws, rr, 2 + j, val, F_IN, "#,##0", hi)
        if v.strip() == "<1":
            c.comment = Comment("Less than one hour a year (deck shows '<1').", "CLASP")
note(ws, 5 + len(t5) - 1, "'43 °C or more' includes 44 °C and above. <1 = less than one hour a year. Source: deck slide 5.")
widths(ws, [34, 12, 16, 16, 16, 16])

# ---------------------------------------------------------------- Hours_40C_and_Above
ws = wb.create_sheet("Hours_40C_and_Above")
title(ws, "Hours at 40 °C and above, degree by degree",
      "National average year (A, B, today) and 10-year counts by zone; stations reaching 44 °C. Links to the other sheets.")
header(ws, 4, ["Bin °C", "Today: hours a year", "A: hours a year", "B: hours a year", "Today: of 1,600", "A: of 1,600", "B: of 1,600"])
HOT = ["40", "41", "42", "43", "44+"]
for k, b in enumerate(HOT):
    rr = 5 + k
    src = FIRST + BINS.index(b)
    put(ws, rr, 1, b if b == "44+" else int(b), F_TXT, bold=True)
    for col, L, fmt in [(2, "B", "#,##0"), (3, "D", "#,##0"), (4, "F", "#,##0"), (5, "C", "0.0"), (6, "E", "0.0"), (7, "G", "0.0")]:
        put(ws, rr, col, f"=Proposed_vs_Current!{L}{src}", F_LINK, fmt)
rr = 5 + len(HOT)
put(ws, rr, 1, "40 °C and above", F_TXT, bold=True, fill=FILL_HI)
for col, fmt in [(2, "#,##0"), (3, "#,##0"), (4, "#,##0"), (5, "0.0"), (6, "0.0"), (7, "0.0")]:
    L = get_column_letter(col)
    put(ws, rr, col, f"=SUM({L}5:{L}{rr - 1})", F_CALC, fmt, FILL_HI)
note(ws, rr + 1, "Today has no 44 °C+ bin ('not counted'). Hours a year are rounded in the deck.")

r0 = rr + 3
header(ws, r0, ["Bin °C: hours counted 2016–2025"] + ZONES + ["All 129 stations"])
for k, b in enumerate(HOT):
    r2 = r0 + 1 + k
    src = B0 + BINS.index(b)
    put(ws, r2, 1, b if b == "44+" else int(b), F_TXT, bold=True)
    for col in range(2, 9):
        L = get_column_letter(col)
        put(ws, r2, col, f"=Zone_Bins_10yr!{L}{src}", F_LINK, "#,##0")
r2 += 1
put(ws, r2, 1, "40 °C and above", F_TXT, bold=True, fill=FILL_HI)
for col in range(2, 9):
    L = get_column_letter(col)
    put(ws, r2, col, f"=SUM({L}{r0 + 1}:{L}{r2 - 1})", F_CALC, "#,##0", FILL_HI)
r2 += 1
put(ws, r2, 1, "Share of the zone's hours at 24 °C+", F_TXT, bold=True)
for col in range(2, 9):
    L = get_column_letter(col)
    put(ws, r2, col, f"={L}{r2 - 1}/Zone_Bins_10yr!{L}{TOT}", F_LINK, "0.00%")

r3 = r2 + 3
header(ws, r3, ["Stations reaching 44 °C or more", "Stations", "Of all stations"])
put(ws, r3 + 1, 1, "At least one hour a year at 44 °C+ (average year)", F_TXT, bold=True)
put(ws, r3 + 1, 2, 50, F_IN, "0")
put(ws, r3 + 1, 3, f"=B{r3 + 1}/B{r3 + 3}", F_CALC, "0.0%")
put(ws, r3 + 2, 1, "Never one hour a year at 44 °C+", F_TXT, bold=True)
put(ws, r3 + 2, 2, f"=B{r3 + 3}-B{r3 + 1}", F_CALC, "0")
put(ws, r3 + 2, 3, f"=B{r3 + 2}/B{r3 + 3}", F_CALC, "0.0%")
put(ws, r3 + 3, 1, "Stations used", F_TXT, bold=True)
put(ws, r3 + 3, 2, 129, F_IN, "0")
note(ws, r3 + 4, "Source: deck slide 9 ('50 of 129 stations reach 44 °C or more for at least an hour in an average year').")

r4 = r3 + 6
header(ws, r4, ["The 8 stations with most hours at 44 °C+", "Zone", "Hours a year at 44 °C+"])
c8 = d["charts"]["s8_c1"]
hotdry = {"Churu", "Sri Ganganagar", "Bikaner", "Jaisalmer", "Barmer"}
for k, (stn, v) in enumerate(zip(c8["categories"], c8["series"]["v"])):
    put(ws, r4 + 1 + k, 1, stn, F_TXT, bold=True)
    put(ws, r4 + 1 + k, 2, "Hot-Dry" if stn in hotdry else "Composite", F_IN)
    put(ws, r4 + 1 + k, 3, v, F_IN, "0")
note(ws, r4 + 9, "Average hours in one year (2016–2025 data). Zone from the bar colour on deck slide 9.")
r5 = r4 + 11
header(ws, r5, ["Not in the deck: needs station-level data"])
for k, t in enumerate([
    "Highest temperature reached at each station (and in any station)",
    "How many stations reach 40, 41, 42 … 45, 46 °C and above",
    "Hours at each degree above 44 °C (45, 46, 47 …), by station",
]):
    put(ws, r5 + 1 + k, 1, t, F_TXT)
widths(ws, [48, 16, 16, 16, 16, 16, 16, 18])

# ---------------------------------------------------------------- Day_Night
ws = wb.create_sheet("Day_Night")
title(ws, "When the warm hours happen: day and night, India (A: by population), average year",
      "Each value = hours a year (days a year that this hour is in the range). Night = 6:30 pm to 5:30 am; day = 6:30 am to 5:30 pm.")
h = d["charts"]["s7_c2"]
night = {"6:30 pm", "7:30 pm", "8:30 pm", "9:30 pm", "10:30 pm", "11:30 pm",
         "12:30 am", "1:30 am", "2:30 am", "3:30 am", "4:30 am", "5:30 am"}
header(ws, 4, ["Hour (IST)", "Day or night", "Hours at 24–34 °C", "Hours at 35 °C+", "Hours at 24 °C+"])
for k, (cat, lo, hi) in enumerate(zip(h["categories"], h["series"]["24–34 °C"], h["series"]["35 °C or more"])):
    rr = 5 + k
    put(ws, rr, 1, cat, F_TXT, bold=True)
    put(ws, rr, 2, "Night" if cat in night else "Day", F_IN)
    put(ws, rr, 3, lo, F_IN, "#,##0.0")
    put(ws, rr, 4, hi, F_IN, "#,##0.0")
    put(ws, rr, 5, f"=C{rr}+D{rr}", F_CALC, "#,##0.0")
H1 = 5 + len(h["categories"]) - 1
rr = H1 + 1
put(ws, rr, 1, "All 24 hours", F_TXT, bold=True)
for col in (3, 4, 5):
    L = get_column_letter(col)
    put(ws, rr, col, f"=SUM({L}5:{L}{H1})", F_CALC, "#,##0.0", FILL_LT)
ALL = rr
s0 = rr + 2
header(ws, s0, ["Summary", "Hours at 24 °C+", "Share of 24 °C+", "Hours at 35 °C+", "Share of 35 °C+"])
for k, per in enumerate(["Day", "Night"]):
    r2 = s0 + 1 + k
    put(ws, r2, 1, f"{per} (12 hours)", F_TXT, bold=True)
    put(ws, r2, 2, f'=SUMIF($B$5:$B${H1},"{per}",$E$5:$E${H1})', F_CALC, "#,##0")
    put(ws, r2, 3, f"=B{r2}/$E${ALL}", F_CALC, "0.0%", FILL_HI if per == "Night" else None)
    put(ws, r2, 4, f'=SUMIF($B$5:$B${H1},"{per}",$D$5:$D${H1})', F_CALC, "#,##0")
    put(ws, r2, 5, f"=D{r2}/$D${ALL}", F_CALC, "0.0%", FILL_HI if per == "Night" else None)
note(ws, s0 + 3, "Deck slide 8: night holds 44% of the hours at 24 °C or more, but only 14% of the hours at 35 °C or more.")

m0 = s0 + 5
mo = d["charts"]["s7_c1"]
header(ws, m0, ["Month", "Hours at 24–34 °C", "Hours at 35 °C+", "Hours at 24 °C+", "In March–October?"])
for k, (cat, lo, hi) in enumerate(zip(mo["categories"], mo["series"]["24–34 °C"], mo["series"]["35 °C or more"])):
    r2 = m0 + 1 + k
    put(ws, r2, 1, cat, F_TXT, bold=True)
    put(ws, r2, 2, lo, F_IN, "#,##0.0")
    put(ws, r2, 3, hi, F_IN, "#,##0.0")
    put(ws, r2, 4, f"=B{r2}+C{r2}", F_CALC, "#,##0.0")
    put(ws, r2, 5, "Yes" if cat in ("Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct") else "No", F_IN)
M1 = m0 + 12
r2 = M1 + 1
put(ws, r2, 1, "Year", F_TXT, bold=True)
for col in (2, 3, 4):
    L = get_column_letter(col)
    put(ws, r2, col, f"=SUM({L}{m0 + 1}:{L}{M1})", F_CALC, "#,##0.0", FILL_LT)
r2 += 1
put(ws, r2, 1, "Share in March–October", F_TXT, bold=True)
put(ws, r2, 4, f'=SUMIF(E{m0 + 1}:E{M1},"Yes",D{m0 + 1}:D{M1})/D{r2 - 1}', F_CALC, "0.0%", FILL_HI)
put(ws, r2, 3, f'=SUMIF(E{m0 + 1}:E{M1},"Yes",C{m0 + 1}:C{M1})/C{r2 - 1}', F_CALC, "0.0%", FILL_HI)
note(ws, r2 + 1, "Deck slide 8: 85% of the hours at 24 °C or more fall in March–October. Per-station day/night needs station-level data.")
widths(ws, [22, 16, 18, 18, 18])

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
wb.save(OUT)
print("wrote", OUT)
